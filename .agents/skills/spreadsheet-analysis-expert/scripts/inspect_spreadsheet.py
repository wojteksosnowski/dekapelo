#!/usr/bin/env python3
"""inspect_spreadsheet.py

Szybkie narzędzie CLI do bezkosztowej inspekcji struktury plików arkuszy kalkulacyjnych (.xlsx, .xls, .ods).
Wypisuje metadane, nazwy arkuszy, wymiary, scalone komórki, potencjalne nagłówki i próbki wierszy
bez zaśmiecania kontekstu LLM zrzutami całych tabel.
"""

import sys
import os
import json
import argparse
from pathlib import Path
from typing import Dict, Any, List, Optional

# Zapewnienie poprawnego kodowania na konsoli Windows
if sys.stdout.encoding and sys.stdout.encoding.lower() != "utf-8":
    try:
        sys.stdout.reconfigure(encoding="utf-8", errors="replace")
    except Exception:
        pass


def inspect_xlsx(filepath: str, sample_rows: int = 5, target_sheet: Optional[str] = None) -> Dict[str, Any]:
    import openpyxl

    wb = openpyxl.load_workbook(filepath, data_only=False)
    # Odczytujemy też wersję z wartościami
    try:
        wb_val = openpyxl.load_workbook(filepath, data_only=True)
    except Exception:
        wb_val = wb

    sheets_info = []
    sheet_names_to_scan = [target_sheet] if target_sheet and target_sheet in wb.sheetnames else wb.sheetnames

    for name in sheet_names_to_scan:
        ws = wb[name]
        ws_val = wb_val[name] if name in wb_val.sheetnames else ws
        
        is_hidden = getattr(ws, "sheet_state", "visible") != "visible"
        max_r = ws.max_row or 0
        max_c = ws.max_column or 0
        
        # Scalone komórki
        merged_ranges = [str(rng) for rng in ws.merged_cells.ranges]
        
        # Wykrycie formuł i próbka
        has_formulas = False
        preview_rows = []
        actual_non_empty_rows = 0
        
        for r_idx in range(1, min(max_r + 1, sample_rows + 15)):
            row_raw = [ws.cell(r_idx, c).value for c in range(1, min(max_c + 1, 20))]
            row_evaluated = [ws_val.cell(r_idx, c).value for c in range(1, min(max_c + 1, 20))]
            
            if any(v is not None for v in row_raw):
                actual_non_empty_rows += 1
            
            for v in row_raw:
                if isinstance(v, str) and v.startswith("="):
                    has_formulas = True
            
            if len(preview_rows) < sample_rows and any(v is not None for v in row_raw):
                row_str = [str(v) if v is not None else "" for v in row_evaluated]
                preview_rows.append(row_str)

        sheets_info.append({
            "sheet_name": name,
            "hidden": is_hidden,
            "max_row": max_r,
            "max_column": max_c,
            "merged_ranges_count": len(merged_ranges),
            "merged_ranges_sample": merged_ranges[:10],
            "has_formulas": has_formulas,
            "sample_rows": preview_rows
        })

    wb.close()
    if wb_val != wb:
        wb_val.close()

    return {
        "file": str(filepath),
        "format": "xlsx",
        "total_sheets": len(wb.sheetnames),
        "sheets": sheets_info
    }


def inspect_via_pandas(filepath: str, ext: str, sample_rows: int = 5, target_sheet: Optional[str] = None) -> Dict[str, Any]:
    import pandas as pd

    engine = "odf" if ext == ".ods" else None
    excel_file = pd.ExcelFile(filepath, engine=engine)
    sheet_names = excel_file.sheet_names
    sheets_info = []

    sheet_names_to_scan = [target_sheet] if target_sheet and target_sheet in sheet_names else sheet_names

    for name in sheet_names_to_scan:
        df = excel_file.parse(name, header=None, nrows=sample_rows + 10)
        sample = []
        for _, row in df.head(sample_rows).iterrows():
            sample.append([str(x) if pd.notna(x) else "" for x in row.tolist()[:15]])

        sheets_info.append({
            "sheet_name": name,
            "hidden": False,
            "max_row": len(df),
            "max_column": df.shape[1] if not df.empty else 0,
            "merged_ranges_count": 0,
            "merged_ranges_sample": [],
            "has_formulas": False,
            "sample_rows": sample
        })

    return {
        "file": str(filepath),
        "format": ext.lstrip("."),
        "total_sheets": len(sheet_names),
        "sheets": sheets_info
    }


def format_text_report(info: Dict[str, Any]) -> str:
    lines = []
    lines.append("=" * 70)
    lines.append(f"[RAPORT INSPEKCJI ARKUSZA] {Path(info['file']).name} [{info['format'].upper()}]")
    lines.append(f"Liczba arkuszy: {info['total_sheets']}")
    lines.append("=" * 70)

    for sh in info["sheets"]:
        status = " (UKRYTY)" if sh.get("hidden") else ""
        lines.append(f"\n[-] Arkusz: '{sh['sheet_name']}'{status}")
        lines.append(f"   Wymiary: {sh['max_row']} wierszy x {sh['max_column']} kolumn")
        lines.append(f"   Scalone komórki: {sh['merged_ranges_count']} (próbka: {', '.join(sh['merged_ranges_sample'][:5]) or 'brak'})")
        lines.append(f"   Obecność formuł: {'TAK' if sh['has_formulas'] else 'NIE lub brak wykrycia'}")
        lines.append("   Podgląd wierszy (pierwsze wiersze danych):")
        for idx, row in enumerate(sh["sample_rows"], 1):
            shortened = [c[:25] + ".." if len(c) > 25 else c for c in row if c]
            lines.append(f"     [w{idx}] " + " | ".join(shortened[:8]))

    lines.append("=" * 70)
    return "\n".join(lines)


def main():
    parser = argparse.ArgumentParser(description="Szybka inspekcja arkuszy kalkulacyjnych .xlsx, .xls, .ods")
    parser.add_argument("file", help="Ścieżka do pliku arkusza")
    parser.add_argument("--sheet", help="Konkretna nazwa arkusza do zbadania (opcjonalne)")
    parser.add_argument("--rows", type=int, default=5, help="Liczba wierszy próbki (domyślnie 5)")
    parser.add_argument("--json", action="store_true", help="Wypisz wynik w formacie JSON")

    args = parser.parse_args()

    file_path = Path(args.file)
    if not file_path.exists():
        print(f"Błąd: Plik {file_path} nie istnieje.", file=sys.stderr)
        sys.exit(1)

    ext = file_path.suffix.lower()

    if ext in [".xlsx", ".xlsm", ".xltx"]:
        res = inspect_xlsx(str(file_path), sample_rows=args.rows, target_sheet=args.sheet)
    elif ext in [".xls", ".ods"]:
        res = inspect_via_pandas(str(file_path), ext, sample_rows=args.rows, target_sheet=args.sheet)
    else:
        print(f"Błąd: Nieobsługiwany format '{ext}'. Dozwolone: .xlsx, .xls, .ods", file=sys.stderr)
        sys.exit(1)

    if args.json:
        print(json.dumps(res, indent=2, ensure_ascii=False))
    else:
        print(format_text_report(res))


if __name__ == "__main__":
    main()
